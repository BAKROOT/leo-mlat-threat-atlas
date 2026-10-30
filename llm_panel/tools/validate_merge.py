#!/usr/bin/env python3
"""
SATERA LLM panel: validate the JSON returned by each LLM, write repair prompts for the
files that fail, and merge everything into Panel_Review.xlsx (+ CSVs) with agreement
statistics and a disagreement list against the catalogue.

Usage (from the llm_panel folder):
    python tools/validate_merge.py
Options:
    --responses responses      folder with the saved JSON files
    --prompts   prompts        folder with manifest.json
    --catalogue ../data/catalogue/SATERA_Extended_Threat_Catalogue.xlsx
    --out       results

File names in the responses folder:  <model>__<prompt>__run<N>.json   e.g. claude__B1__run1.json
A correction round for the same record set:  <model>__<prompt>__run<N>__fix1.json
(<model> = claude | chatgpt | deepseek | gemini ; <prompt> = A | B1 | B2 | B3 | B4)
Requires: Python 3.9+, pandas, numpy, scipy, openpyxl.
"""
import argparse, json, math, os, re, sys, itertools, collections
import numpy as np
import pandas as pd
import openpyxl
from scipy.stats import spearmanr

# ------------------------------------------------------------------ enumerations
STRIDE = ["S", "T", "R", "I", "D", "E", "tau"]
GOALS = ["Displacement", "Denial", "Deception", "Disclosure"]
TIERS = ["T1p", "T1c", "T2", "T3", "T3x"]
EVID = ["P", "E0", "E1", "E2", "E3"]
LAYERS = ["L1", "L2", "L3", "L4"]
ACS = [f"AC{i}" for i in range(1, 8)]
TBS = ["TB0", "TB1", "TB2", "TB3"]
VALID = ["valid", "duplicate", "not_a_threat", "out_of_scope"]
SCORES = ["D", "R", "E", "A", "Di_strong", "Di_weak"]
BASE_R = {"P": 1, "E0": 2, "E1": 3, "E2": 4, "E3": 5}
E_OF_TIER = {"T1p": 5, "T1c": 4, "T2": 3, "T3": 2, "T3x": 1}
CANON = {  # lower-case -> canonical
    **{s.lower(): s for s in STRIDE}, "τ": "tau", "temporal": "tau",
    **{g.lower(): g for g in GOALS}, **{t.lower(): t for t in TIERS},
    **{e.lower(): e for e in EVID}, **{l.lower(): l for l in LAYERS},
    **{a.lower(): a for a in ACS}, **{t.lower(): t for t in TBS}, **{v: v for v in VALID},
}
ASSET_RE = re.compile(r"^A[1-4]\.\d$")


def band(s):
    if s is None or (isinstance(s, float) and math.isnan(s)):
        return None
    return "Low" if s < 10 else "Medium" if s < 15 else "High" if s < 20 else "Critical"


# ------------------------------------------------------------------ tolerant JSON loading
def load_json_text(text):
    """Return (obj, note). Accepts code fences, text around the object, // comments, trailing commas."""
    notes = []
    t = text.strip()
    m = re.findall(r"```(?:json|JSON)?\s*(.*?)```", t, flags=re.S)
    if m:
        t = max(m, key=len).strip()
        notes.append("code fence removed")
    i, j = t.find("{"), t.rfind("}")
    if i < 0:
        raise ValueError("no '{' found: the reply contains no JSON object")
    if j < i:
        raise ValueError("JSON is cut off: no closing '}' (output truncated?)")
    t = t[i:j + 1]
    t = t.replace("“", '"').replace("”", '"').replace("‘", "'").replace("’", "'")
    try:
        return json.loads(t), notes
    except json.JSONDecodeError:
        pass
    t2 = re.sub(r'(?<!:)(?<!")//[^\n"]*(?=\n)', "", t)       # // comments at line end
    t2 = re.sub(r",\s*([}\]])", r"\1", t2)                    # trailing commas
    try:
        return json.loads(t2), notes + ["comments/trailing commas repaired"]
    except json.JSONDecodeError as e:
        a, b = max(0, e.pos - 60), e.pos + 60
        depth = t2.count("{") - t2.count("}")
        hint = " (braces unbalanced: output probably truncated)" if depth else ""
        raise ValueError(f"invalid JSON at char {e.pos}: {e.msg}{hint}; near: ...{t2[a:b]!r}...")


def as_list(v):
    if v is None:
        return []
    if isinstance(v, list):
        return v
    if isinstance(v, str):
        return [x.strip() for x in re.split(r"[;,]", v) if x.strip()]
    return [v]


def canon(v):
    if not isinstance(v, str):
        return v
    k = v.strip()
    return CANON.get(k.lower(), k)


def to_int(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)) and float(v) == int(v):
        return int(v)
    if isinstance(v, str) and re.fullmatch(r"\s*\d+\s*", v):
        return int(v)
    return None


def to_bool(v):
    if isinstance(v, bool):
        return v
    if isinstance(v, str) and v.strip().lower() in ("true", "false"):
        return v.strip().lower() == "true"
    return None


# ------------------------------------------------------------------ record validation
def check_rating(rec, expected_layer_prefix, asset_ids):
    """Normalise one rating record; return (clean_record, errors, warnings)."""
    err, warn = [], []
    c = {}
    c["id"] = str(rec.get("id", "")).strip().upper()
    c["validity"] = canon(rec.get("validity"))
    if c["validity"] not in VALID:
        err.append(f"validity={rec.get('validity')!r}")
    rid = rec.get("related_id")
    c["related_id"] = None if rid in (None, "", "null") else str(rid).strip().upper()
    c["layer"] = canon(rec.get("layer"))
    if c["layer"] not in LAYERS:
        err.append(f"layer={rec.get('layer')!r}")
    c["asset_ids"] = [str(a).strip().upper() for a in as_list(rec.get("asset_ids"))]
    if not c["asset_ids"]:
        err.append("asset_ids empty")
    for a in c["asset_ids"]:
        if not ASSET_RE.match(a) or a not in asset_ids:
            err.append(f"unknown asset {a}")
    c["trust_boundary"] = [canon(x) for x in as_list(rec.get("trust_boundary")) if str(x).strip().lower() not in ("none", "")]
    for x in c["trust_boundary"]:
        if x not in TBS:
            err.append(f"trust_boundary {x!r}")
    c["stride_primary"] = canon(rec.get("stride_primary"))
    if c["stride_primary"] not in STRIDE:
        err.append(f"stride_primary={rec.get('stride_primary')!r}")
    c["stride_secondary"] = [canon(x) for x in as_list(rec.get("stride_secondary"))]
    for x in c["stride_secondary"]:
        if x not in STRIDE:
            err.append(f"stride_secondary {x!r}")
    if len(c["stride_secondary"]) > 2:
        warn.append("more than 2 secondary categories")
    c["goals"] = [canon(x) for x in as_list(rec.get("goals"))]
    for x in c["goals"]:
        if x not in GOALS:
            err.append(f"goal {x!r}")
    c["is_enabler"] = to_bool(rec.get("is_enabler"))
    if c["is_enabler"] is None:
        err.append("is_enabler not boolean")
    c["adversary_class"] = [canon(x) for x in as_list(rec.get("adversary_class"))]
    if not c["adversary_class"]:
        err.append("adversary_class empty")
    for x in c["adversary_class"]:
        if x not in ACS:
            err.append(f"adversary_class {x!r}")
    c["tier"] = canon(rec.get("tier"))
    if c["tier"] not in TIERS:
        err.append(f"tier={rec.get('tier')!r}")
    c["evidence"] = canon(rec.get("evidence"))
    if c["evidence"] not in EVID:
        err.append(f"evidence={rec.get('evidence')!r}")
    c["analogue_only"] = to_bool(rec.get("analogue_only"))
    if c["analogue_only"] is None:
        err.append("analogue_only not boolean")
    c["unknown_sources"] = [str(x).lstrip("#") for x in as_list(rec.get("unknown_sources"))]
    for k in SCORES + ["confidence"]:
        v = rec.get(k)
        if k == "D" and (v is None or str(v).lower() == "null"):
            c[k] = None
            continue
        iv = to_int(v)
        if iv is None or not 1 <= iv <= 5:
            err.append(f"{k}={v!r} (need integer 1-5)")
            iv = None
        c[k] = iv
    if c["is_enabler"] is True and c["D"] is not None:
        warn.append("enabler but D given")
    if c["is_enabler"] is False and c["D"] is None and not any("D=" in e for e in err):
        err.append("D is null but is_enabler is false")
    if c["is_enabler"] is True and c["goals"]:
        warn.append("enabler but goals given")
    if c["is_enabler"] is False and not c["goals"]:
        warn.append("no goal given for a non-enabler")
    # internal consistency with the rubric (warnings only)
    if c["evidence"] in BASE_R and c["R"] is not None and c["analogue_only"] is not None:
        exp = max(1, BASE_R[c["evidence"]] - (1 if c["analogue_only"] else 0))
        if c["R"] != exp:
            warn.append(f"R={c['R']} but evidence {c['evidence']}/analogue={c['analogue_only']} gives {exp}")
    if c["tier"] in E_OF_TIER and c["E"] is not None and c["E"] != E_OF_TIER[c["tier"]]:
        warn.append(f"E={c['E']} but tier {c['tier']} gives {E_OF_TIER[c['tier']]}")
    for k in ("rationale", "issue"):
        c[k] = str(rec.get(k, "") or "")
    if c["validity"] in VALID and c["validity"] != "valid" and not c["issue"]:
        warn.append("not valid but no issue text")
    return c, err, warn


def check_elicit(rec, asset_ids):
    err, c = [], {}
    c["title"] = str(rec.get("title", "")).strip()
    if not c["title"]:
        err.append("title empty")
    c["layer"] = canon(rec.get("layer"))
    if c["layer"] not in LAYERS:
        err.append(f"layer={rec.get('layer')!r}")
    c["asset_ids"] = [str(a).strip().upper() for a in as_list(rec.get("asset_ids"))]
    for a in c["asset_ids"]:
        if a not in asset_ids:
            err.append(f"unknown asset {a}")
    c["mechanism"] = str(rec.get("mechanism", ""))
    c["stride_primary"] = canon(rec.get("stride_primary"))
    if c["stride_primary"] not in STRIDE:
        err.append(f"stride_primary={rec.get('stride_primary')!r}")
    c["goals"] = [canon(x) for x in as_list(rec.get("goals"))]
    for x in c["goals"]:
        if x not in GOALS:
            err.append(f"goal {x!r}")
    c["tier"] = canon(rec.get("tier"))
    if c["tier"] not in TIERS:
        err.append(f"tier={rec.get('tier')!r}")
    c["evidence_support"] = canon(rec.get("evidence_support"))
    if c["evidence_support"] not in EVID:
        err.append(f"evidence_support={rec.get('evidence_support')!r}")
    sv = to_int(rec.get("severity"))
    if sv is None or not 1 <= sv <= 5:
        err.append(f"severity={rec.get('severity')!r}")
    c["severity"] = sv
    c["commonly_overlooked"] = to_bool(rec.get("commonly_overlooked"))
    return c, err


# ------------------------------------------------------------------ catalogue (reference answers)
def load_catalogue(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Catalogue"]
    h = {ws.cell(4, c).value: c for c in range(1, ws.max_column + 1)}
    rows = []
    for r in range(5, ws.max_row + 1):
        tid = ws.cell(r, 1).value
        if not tid or not re.match(r"T\d\.\d\d$", str(tid)):
            continue
        def g(k):
            v = ws.cell(r, h[k]).value
            # whole-number scores are stored as 3 or 3.0 depending on how the workbook was saved: read both as int
            return int(v) if isinstance(v, float) and v.is_integer() else v
        rows.append(dict(
            id=tid, layer=f"L{str(tid)[1]}", title=g("Threat"),
            stride_primary=g("STRIDE-tau primary"), stride_secondary=g("Secondary"),
            tier=g("Tier code"), evidence=g("Evidence"), is_enabler=(g("Role") == "Enabler"),
            goals=[x for x in GOALS if g(x) == "Y"], assets=[a.strip() for a in str(g("Asset(s)")).split(",")],
            D=g("D"), R=g("R"), E=g("E"), A=g("A"), Di_strong=g("Di strong"), Di_weak=g("Di weak"),
            S_strong=g("S strong"), S_weak=g("S weak"), band_strong=g("Band strong"), band_weak=g("Band weak"),
            adversary_class=g("Adversary class")))
    wa = wb["Assets"]
    assets = {wa.cell(r, 1).value: wa.cell(r, 2).value for r in range(4, wa.max_row + 1)
              if wa.cell(r, 1).value and ASSET_RE.match(str(wa.cell(r, 1).value))}
    return pd.DataFrame(rows).set_index("id"), assets


# ------------------------------------------------------------------ Krippendorff's alpha
def kalpha(units, metric="nominal"):
    """units: list of lists of values (one list per item, missing values omitted)."""
    o = collections.defaultdict(float)
    for u in units:
        m = len(u)
        if m < 2:
            continue
        for i, j in itertools.permutations(range(m), 2):
            o[(u[i], u[j])] += 1.0 / (m - 1)
    vals = sorted({k[0] for k in o} | {k[1] for k in o}, key=str)
    n = sum(o.values())
    if n <= 1 or len(vals) < 2:
        return float("nan")
    nc = {c: sum(o.get((c, k), 0) for k in vals) for c in vals}
    d = (lambda a, b: 0.0 if a == b else 1.0) if metric == "nominal" else (lambda a, b: float(a - b) ** 2)
    do = sum(o.get((c, k), 0) * d(c, k) for c in vals for k in vals) / n
    de = sum(nc[c] * nc[k] * d(c, k) for c in vals for k in vals) / (n * (n - 1))
    return float("nan") if de == 0 else 1 - do / de


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--responses", default="responses")
    ap.add_argument("--prompts", default="prompts")
    ap.add_argument("--catalogue", default="../data/catalogue/SATERA_Extended_Threat_Catalogue.xlsx")
    ap.add_argument("--out", default="results")
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    manifest = json.load(open(os.path.join(a.prompts, "manifest.json")))
    cat, assets = load_catalogue(a.catalogue)
    asset_ids = set(assets)
    all_expected = [i for p in manifest.values() if p["expected_ids"] and not p.get("small") for i in p["expected_ids"]]
    assert sorted(all_expected) == sorted(cat.index), "manifest and catalogue disagree"

    files = sorted(f for f in os.listdir(a.responses) if f.lower().endswith((".json", ".txt")))
    pat = re.compile(r"^(?P<model>[a-z0-9]+)__(?P<prompt>A|B[1-4][a-h]?)__run(?P<run>\d+)(?:__fix(?P<fix>\d+))?\.(json|txt)$", re.I)
    groups = collections.defaultdict(list)
    report, repairs, bad_names = [], {}, []
    for f in files:
        m = pat.match(f)
        if not m:
            bad_names.append(f)
            continue
        pr = m["prompt"]; pr = "A" if pr.upper() == "A" else pr[:2].upper() + pr[2:].lower()
        key = (m["model"].lower(), pr, int(m["run"]))
        groups[key].append((int(m["fix"] or 0), f))
    for f in bad_names:
        report.append(f"[NAME] {f}: does not match <model>__<prompt>__run<N>[__fix<K>].json, ignored")

    long_rows, elic_rows, val_rows = [], [], []
    for (model, pid, run), flist in sorted(groups.items()):
        flist.sort()
        records, parse_err, notes = {}, [], []
        elic_records = []
        meta_ok = True
        for fix, f in flist:
            try:
                text = open(os.path.join(a.responses, f), encoding="utf-8").read()
                obj, nt = load_json_text(text)
                notes += [f"{f}: {x}" for x in nt]
            except Exception as e:
                parse_err.append(f"{f}: {e}")
                continue
            if pid == "A":
                lst = obj.get("threats")
                if not isinstance(lst, list):
                    parse_err.append(f"{f}: key 'threats' missing or not a list")
                else:
                    elic_records = lst
                continue
            lst = obj.get("ratings")
            if not isinstance(lst, list):
                parse_err.append(f"{f}: key 'ratings' missing or not a list")
                continue
            for rec in lst:
                if isinstance(rec, dict):
                    records[str(rec.get("id", "")).strip().upper()] = rec   # fix files override
        tag = f"{model} {pid} run{run}"
        status, errs, warns = "OK", [], []
        if pid == "A":
            if parse_err:
                status = "PARSE ERROR"; errs = parse_err
            else:
                try:
                    declared = obj.get("meta", {}).get("n_threats")
                    if declared is not None and declared != len(elic_records):
                        warns.append(f"meta.n_threats={declared} but {len(elic_records)} threats present")
                except Exception:
                    pass
                if len(elic_records) > 45:
                    warns.append(f"{len(elic_records)} threats (limit 45)")
                for i, rec in enumerate(elic_records):
                    c, e = check_elicit(rec, asset_ids)
                    if e:
                        errs.append(f"threat #{i+1} '{c['title']}': " + "; ".join(e))
                    c.update(model=model, run=run)
                    elic_rows.append(c)
                if errs:
                    status = "FIELD ERRORS"
            if status != "OK" and errs:
                repairs[(model, pid, run)] = (status, errs, None)
        else:
            expected = manifest[pid]["expected_ids"]
            if parse_err and not records:
                status = "PARSE ERROR"; errs = parse_err
                repairs[(model, pid, run)] = (status, errs, None)
            else:
                warns += [f"(earlier file superseded or unreadable) {x}" for x in parse_err]
                missing = [i for i in expected if i not in records]
                extra = [i for i in records if i not in expected]
                bad_ids = []
                for tid in expected:
                    if tid not in records:
                        continue
                    c, e, w = check_rating(records[tid], None, asset_ids)
                    if e:
                        errs.append(f"{tid}: " + "; ".join(e)); bad_ids.append(tid)
                    for x in w:
                        warns.append(f"{tid}: {x}")
                    c.update(model=model, run=run, prompt=pid, field_errors=bool(e))
                    long_rows.append(c)
                if missing:
                    errs.append("missing ids: " + ", ".join(missing))
                if extra:
                    warns.append("unexpected ids ignored: " + ", ".join(extra))
                if errs:
                    status = "ERRORS"
                    repairs[(model, pid, run)] = (status, errs, sorted(set(missing + bad_ids)))
        val_rows.append(dict(model=model, prompt=pid, run=run, files=", ".join(f for _, f in flist), status=status,
                             n_errors=len(errs), n_warnings=len(warns)))
        report.append(f"\n=== {tag}: {status}  ({', '.join(f for _, f in flist)})")
        for n in notes:
            report.append(f"  [auto-fix] {n}")
        for e in errs:
            report.append(f"  [ERROR] {e}")
        for w in warns[:40]:
            report.append(f"  [warn] {w}")
        if len(warns) > 40:
            report.append(f"  [warn] ... {len(warns)-40} more warnings")

    # coverage matrix: which model/prompt files are present
    models = sorted({k[0] for k in groups})
    cov = ["\n=== COVERAGE (status of run 1 per file set; ids covered = threats with a usable rating)"]
    ok_ids = collections.defaultdict(set)
    for r in long_rows:
        if not r.get("field_errors"):
            ok_ids[r["model"]].add(r["id"])
    for m in models:
        parts = [f"{p}:{v['status']}" for p in manifest for v in val_rows if v["model"] == m and v["prompt"] == p and v["run"] == 1]
        cov.append("  " + m.ljust(10) + "  ".join(parts) + f"   | threats rated: {len(ok_ids[m])}/55")
    report = cov + report

    # repair prompts
    rep_dir = os.path.join(a.out, "repair_prompts")
    os.makedirs(rep_dir, exist_ok=True)
    for f in os.listdir(rep_dir):
        os.remove(os.path.join(rep_dir, f))
    for (model, pid, run), (status, errs, ids) in repairs.items():
        name = f"repair_{model}__{pid}__run{run}.txt"
        if status == "PARSE ERROR":
            body = ("Your previous reply could not be parsed as JSON.\nProblem: " + " | ".join(errs)[:900] +
                    "\n\nReturn the COMPLETE JSON object again, in ONE code block, valid JSON only (double quotes, no comments, "
                    "no trailing commas, no line breaks inside strings). If the reply was cut off, shorten the free-text fields "
                    "(rationale, issue) so that every record fits. Do not add any text outside the code block.")
        elif pid == "A":
            body = ("Some threats in your JSON have invalid field values:\n- " + "\n- ".join(errs[:30]) +
                    "\n\nReturn ONE JSON object in ONE code block with the same schema, where \"threats\" contains ONLY the corrected "
                    "versions of the threats listed above (same titles). Valid JSON only; nothing outside the code block.")
        else:
            body = ("Some records in your JSON are missing or have invalid values:\n- " + "\n- ".join(errs[:40]) +
                    f"\n\nReturn ONE JSON object in ONE code block with the same schema (keep \"meta\"), where \"ratings\" contains ONLY "
                    f"corrected records for these ids: {', '.join(ids)}. Use exactly the allowed values of the original format. "
                    "Valid JSON only; nothing outside the code block.\n(Save your reply as "
                    f"{model}__{pid}__run{run}__fix1.json)")
        open(os.path.join(rep_dir, name), "w", encoding="utf-8").write(body)

    open(os.path.join(a.out, "validation_report.txt"), "w", encoding="utf-8").write("\n".join(report))
    print("\n".join(report))

    def run_elicit(E_df):
        stop = set("the a an of to in on and or for with by from that this which is are be it its as at into than then their can may not no via any all each one two both after before so such such".split())
        toks = lambda s: {w for w in re.findall(r"[a-z0-9]{4,}", str(s).lower()) if w not in stop}
        match_rows, gap_rows = [], []
        if len(E_df):
            for m in sorted(E_df.model.unique()):
                em = E_df[E_df.model == m].reset_index(drop=True)
                best_for_model = []
                for tid, ref in cat.iterrows():
                    ct = toks(ref.title) | toks(ref.get("title", ""))
                    best, bs = None, 0
                    for i, e in em.iterrows():
                        tj = len(toks(e.title + " " + e.mechanism) & ct) / max(1, len(ct))
                        aj = len(set(e.asset_ids) & set(ref.assets)) / max(1, len(set(e.asset_ids) | set(ref.assets)))
                        sc = 0.45 * min(1, tj * 1.5) + 0.40 * aj + 0.15 * (e.layer == ref.layer)
                        if sc > bs:
                            best, bs = i, sc
                    match_rows.append(dict(model=m, catalogue_id=tid, catalogue_title=ref.title,
                                           best_model_threat=em.title[best] if best is not None else "", score=round(bs, 2),
                                           likely_match=bs >= 0.45))
                for i, e in em.iterrows():
                    bs, bid = 0, None
                    for tid, ref in cat.iterrows():
                        ct = toks(ref.title)
                        tj = len(toks(e.title + " " + e.mechanism) & ct) / max(1, len(ct))
                        aj = len(set(e.asset_ids) & set(ref.assets)) / max(1, len(set(e.asset_ids) | set(ref.assets)))
                        sc = 0.45 * min(1, tj * 1.5) + 0.40 * aj + 0.15 * (e.layer == ref.layer)
                        if sc > bs:
                            bs, bid = sc, tid
                    if bs < 0.45:
                        gap_rows.append(dict(model=m, title=e.title, layer=e.layer, assets=";".join(e.asset_ids), stride=e.stride_primary,
                                             severity=e.severity, commonly_overlooked=e.commonly_overlooked, mechanism=e.mechanism,
                                             closest_catalogue_id=bid, closest_score=round(bs, 2)))
        MR, GR = pd.DataFrame(match_rows), pd.DataFrame(gap_rows)
        return MR, GR

    if not long_rows:
        print("\nNo rating files parsed yet: nothing to merge.")
        E_df = pd.DataFrame(elic_rows)
        if len(E_df):
            MR, GR = run_elicit(E_df)
            xp = os.path.join(a.out, "Elicitation_Review.xlsx")
            with pd.ExcelWriter(xp, engine="openpyxl") as xw:
                E_df.assign(asset_ids=E_df.asset_ids.map(";".join), goals=E_df.goals.map(";".join)).to_excel(xw, sheet_name="Elicited_all", index=False)
                MR.to_excel(xw, sheet_name="Elicit_coverage_heur", index=False)
                GR.to_excel(xw, sheet_name="Elicit_candidate_gaps", index=False)
            print(f"Wrote {xp}")
        return

    # ------------------------------------------------------------------ merge ratings
    L = pd.DataFrame(long_rows)
    L = L[~L.field_errors].copy()
    L["S_strong"] = np.where(L.D.notna() & ~L.is_enabler.astype(bool), L.D + L.R + L.E + L.A + L.Di_strong, np.nan)
    L["S_weak"] = np.where(L.D.notna() & ~L.is_enabler.astype(bool), L.D + L.R + L.E + L.A + L.Di_weak, np.nan)
    L["band_strong"] = L.S_strong.map(band)
    L["band_weak"] = L.S_weak.map(band)
    for c in ("asset_ids", "trust_boundary", "stride_secondary", "goals", "adversary_class", "unknown_sources"):
        L[c + "_txt"] = L[c].map(lambda x: ";".join(x))
    for g in GOALS:
        L["goal_" + g] = L.goals.map(lambda x: g in x).astype(int)

    # one row per model per threat (collapse runs: median of scores, mode of categories)
    def collapse(d):
        r = {}
        for k in ["validity", "layer", "stride_primary", "tier", "evidence", "is_enabler", "analogue_only"]:
            r[k] = d[k].mode().iloc[0]
        for k in SCORES + ["confidence", "S_strong", "S_weak"]:
            r[k] = d[k].median()
        for g in GOALS:
            r["goal_" + g] = int(d["goal_" + g].mean() >= 0.5)
        r["asset_ids_txt"] = d["asset_ids_txt"].iloc[0]
        r["stride_secondary_txt"] = d["stride_secondary_txt"].iloc[0]
        r["adversary_class_txt"] = d["adversary_class_txt"].iloc[0]
        r["stride_any"] = d["stride_primary"].tolist() + sum(d["stride_secondary"].tolist(), [])
        r["issue"] = " | ".join(x for x in d["issue"] if x)[:300]
        r["rationale"] = d["rationale"].iloc[0]
        r["n_runs"] = len(d)
        return pd.Series(r)

    M = L.groupby(["model", "id"]).apply(collapse, include_groups=False).reset_index()
    M["band_strong"] = M.S_strong.map(band)
    M["band_weak"] = M.S_weak.map(band)
    models = sorted(M.model.unique())

    # per-threat summary
    srows, drows = [], []
    for tid, ref in cat.iterrows():
        d = M[M.id == tid]
        n = len(d)
        if n == 0:
            continue
        row = dict(id=tid, title=ref.title, n_models=n)
        flags = []

        def cat_field(name, refval):
            vc = d[name].astype(str).value_counts()
            top, share = vc.index[0], vc.iloc[0] / n
            cons = share >= 2 / 3
            match = str(top) == str(refval)
            t = "agree" if cons and match else "panel_vs_catalogue" if cons else "split"
            row[f"{name}_panel"], row[f"{name}_cat"], row[f"{name}_share"], row[f"{name}_tag"] = top, refval, round(share, 2), t
            return t

        def score_field(name, refval):
            v = d[name].dropna()
            if len(v) == 0 or pd.isna(refval):
                return None
            med = float(v.median()); iqr = float(np.percentile(v, 75) - np.percentile(v, 25))
            cons = iqr <= 1
            match = abs(med - float(refval)) <= 1
            t = "agree" if cons and match else "panel_vs_catalogue" if cons else "split"
            row[f"{name}_panel_median"], row[f"{name}_cat"], row[f"{name}_iqr"], row[f"{name}_tag"] = med, refval, round(iqr, 2), t
            return t

        tags = {}
        for f_, rv in [("layer", ref.layer), ("stride_primary", ref.stride_primary), ("tier", ref.tier),
                       ("evidence", ref.evidence), ("is_enabler", bool(ref.is_enabler)), ("validity", "valid")]:
            tags[f_] = cat_field(f_, rv)
        for f_ in SCORES:
            t = score_field(f_, ref[f_])
            if t:
                tags[f_] = t
        # goals: per-goal vote
        gtag = "agree"
        for g in GOALS:
            share = d["goal_" + g].mean()
            inref = g in ref.goals
            row[f"goal_{g}_share"] = round(share, 2)
            if 1 / 3 < share < 2 / 3:
                gtag = "split" if gtag != "panel_vs_catalogue" else gtag
            elif (share >= 2 / 3) != inref and gtag == "agree":
                gtag = "panel_vs_catalogue"
        row["goals_cat"] = ";".join(ref.goals)
        row["goals_panel"] = ";".join(g for g in GOALS if d["goal_" + g].mean() > 0.5)
        tags["goals"] = gtag
        s = d.S_strong.dropna()
        row["S_strong_panel_median"] = float(s.median()) if len(s) else np.nan
        row["S_strong_cat"] = ref.S_strong
        row["band_strong_panel"] = band(row["S_strong_panel_median"]) if len(s) else None
        row["band_strong_cat"] = ref.band_strong
        s2 = d.S_weak.dropna()
        row["S_weak_panel_median"] = float(s2.median()) if len(s2) else np.nan
        row["S_weak_cat"] = ref.S_weak
        row["stride_cat_in_any_panel_reply"] = int(sum(ref.stride_primary in x for x in d.stride_any))
        row["n_flagged_not_valid"] = int((d.validity != "valid").sum())
        row["issues_raised"] = " || ".join(f"{m}: {x}" for m, x in zip(d.model, d.issue) if x)[:500]
        n_pvc = sum(v == "panel_vs_catalogue" for v in tags.values())
        n_split = sum(v == "split" for v in tags.values())
        row["fields_agree"] = sum(v == "agree" for v in tags.values())
        row["fields_panel_vs_catalogue"] = n_pvc
        row["fields_split"] = n_split
        row["needs_adjudication"] = bool(n_pvc or n_split or row["n_flagged_not_valid"])
        for k, v in tags.items():
            row[f"tag_{k}"] = v
            if v != "agree":
                drows.append(dict(id=tid, title=ref.title, field=k, tag=v,
                                  catalogue=row.get(f"{k}_cat", row.get("goals_cat") if k == "goals" else ""),
                                  panel=row.get(f"{k}_panel", row.get(f"{k}_panel_median", row.get("goals_panel") if k == "goals" else "")),
                                  spread=row.get(f"{k}_share", row.get(f"{k}_iqr", ""))))
        srows.append(row)
    S = pd.DataFrame(srows)
    Dg = pd.DataFrame(drows)

    # ------------------------------------------------------------------ agreement statistics
    ag = []
    wide = lambda col: M.pivot(index="id", columns="model", values=col)
    for col in ["layer", "stride_primary", "tier", "evidence", "is_enabler", "validity"]:
        w = wide(col)
        ag.append(dict(metric=f"Krippendorff alpha (nominal): {col}", value=kalpha([[str(x) for x in r.dropna()] for _, r in w.iterrows()]), n_items=len(w)))
    for g in GOALS:
        w = wide("goal_" + g)
        ag.append(dict(metric=f"Krippendorff alpha (nominal): goal {g}", value=kalpha([[int(x) for x in r.dropna()] for _, r in w.iterrows()]), n_items=len(w)))
    for col in SCORES + ["S_strong", "S_weak"]:
        w = wide(col)
        ag.append(dict(metric=f"Krippendorff alpha (interval): {col}", value=kalpha([[float(x) for x in r.dropna()] for _, r in w.iterrows()], "interval"), n_items=len(w)))
    A_ = pd.DataFrame(ag)

    per_model = []
    for m in models:
        d = M[M.model == m].set_index("id").join(cat, rsuffix="_cat")
        r = dict(model=m, n_threats=len(d))
        for k in ["layer", "stride_primary", "tier", "evidence", "is_enabler"]:
            r[f"{k} exact match"] = round(float((d[k].astype(str) == d[k + "_cat"].astype(str)).mean()), 3)
        r["stride: catalogue primary in model primary or secondary"] = round(float(np.mean([cat.loc[i, "stride_primary"] in x for i, x in d.stride_any.items()])), 3)
        for k in SCORES:
            ok = d[k].notna() & d[k + "_cat"].notna()
            r[f"{k} within +-1"] = round(float((abs(d.loc[ok, k] - d.loc[ok, k + "_cat"]) <= 1).mean()), 3) if ok.any() else np.nan
            r[f"{k} MAE"] = round(float(abs(d.loc[ok, k] - d.loc[ok, k + "_cat"]).mean()), 2) if ok.any() else np.nan
        jac = []
        for i, row in d.iterrows():
            a_ = {g for g in GOALS if row["goal_" + g]}; b_ = set(cat.loc[i, "goals"])
            jac.append(1.0 if not a_ and not b_ else len(a_ & b_) / len(a_ | b_))
        r["goal-set Jaccard (mean)"] = round(float(np.mean(jac)), 3)
        ok = d.S_strong.notna() & d.S_strong_cat.notna()
        if ok.sum() > 2:
            r["S_strong Spearman vs catalogue"] = round(float(spearmanr(d.loc[ok, "S_strong"], d.loc[ok, "S_strong_cat"])[0]), 3)
            ok2 = ok & d.band_strong.notna()
            r["band_strong exact match"] = round(float((d.loc[ok2, "band_strong"] == d.loc[ok2, "band_strong_cat"]).mean()), 3)
            top = lambda s_: set(s_.sort_values(ascending=False).head(10).index)
            r["top-10 overlap (S_strong)"] = len(top(d.loc[ok, "S_strong"]) & top(d.loc[ok, "S_strong_cat"]))
        r["share valid"] = round(float((d.validity == "valid").mean()), 3)
        r["mean confidence"] = round(float(d.confidence.mean()), 2)
        per_model.append(r)
    PM = pd.DataFrame(per_model)
    pw = []
    for m1, m2 in itertools.combinations(models, 2):
        w = wide("S_strong")[[m1, m2]].dropna()
        pw.append(dict(model_a=m1, model_b=m2, n=len(w),
                       S_strong_spearman=round(float(spearmanr(w[m1], w[m2])[0]), 3) if len(w) > 2 else np.nan,
                       stride_primary_agree=round(float((wide("stride_primary")[[m1, m2]].dropna().nunique(axis=1) == 1).mean()), 3),
                       tier_agree=round(float((wide("tier")[[m1, m2]].dropna().nunique(axis=1) == 1).mean()), 3)))
    PW = pd.DataFrame(pw)
    summ = pd.DataFrame([
        dict(metric="threats rated by at least one model", value=len(S)),
        dict(metric="threats needing adjudication", value=int(S.needs_adjudication.sum()) if len(S) else 0),
        dict(metric="models in panel", value=len(models)),
        dict(metric="panel S_strong Spearman vs catalogue (median of models)",
             value=round(float(spearmanr(S.dropna(subset=["S_strong_panel_median", "S_strong_cat"]).S_strong_panel_median,
                                         S.dropna(subset=["S_strong_panel_median", "S_strong_cat"]).S_strong_cat.astype(float))[0]), 3)
             if S[["S_strong_panel_median", "S_strong_cat"]].dropna().shape[0] > 2 else np.nan),
        dict(metric="panel band_strong equals catalogue band (share of scored threats)",
             value=round(float((S.dropna(subset=["band_strong_panel"]).band_strong_panel == S.dropna(subset=["band_strong_panel"]).band_strong_cat).mean()), 3)
             if S.band_strong_panel.notna().any() else np.nan),
    ])

    E_df = pd.DataFrame(elic_rows)
    MR, GR = run_elicit(E_df)

    # ------------------------------------------------------------------ write
    keep_long = ["model", "id", "validity", "layer", "asset_ids_txt", "stride_primary", "stride_secondary_txt", "goal_Displacement",
                 "goal_Denial", "goal_Deception", "goal_Disclosure", "is_enabler", "adversary_class_txt", "tier", "evidence", "analogue_only",
                 "D", "R", "E", "A", "Di_strong", "Di_weak", "S_strong", "band_strong", "S_weak", "band_weak", "confidence", "issue", "rationale", "n_runs"]
    M[keep_long].to_csv(os.path.join(a.out, "panel_long.csv"), index=False)
    S.to_csv(os.path.join(a.out, "panel_summary.csv"), index=False)
    Dg.to_csv(os.path.join(a.out, "disagreements.csv"), index=False)
    PM.to_csv(os.path.join(a.out, "model_vs_catalogue.csv"), index=False)
    A_.to_csv(os.path.join(a.out, "agreement_alpha.csv"), index=False)
    xp = os.path.join(a.out, "Panel_Review.xlsx")
    with pd.ExcelWriter(xp, engine="openpyxl") as xw:
        summ.to_excel(xw, sheet_name="Overview", index=False)
        A_.to_excel(xw, sheet_name="Agreement", index=False)
        PM.to_excel(xw, sheet_name="Model_vs_catalogue", index=False)
        PW.to_excel(xw, sheet_name="Model_pairs", index=False)
        S.to_excel(xw, sheet_name="Per_threat", index=False)
        Dg.to_excel(xw, sheet_name="Disagreements", index=False)
        M[keep_long].to_excel(xw, sheet_name="Long", index=False)
        pd.DataFrame(val_rows).to_excel(xw, sheet_name="Validation", index=False)
        if len(E_df):
            E_df.assign(asset_ids=E_df.asset_ids.map(";".join), goals=E_df.goals.map(";".join)).to_excel(xw, sheet_name="Elicited_all", index=False)
            MR.to_excel(xw, sheet_name="Elicit_coverage_heur", index=False)
            GR.to_excel(xw, sheet_name="Elicit_candidate_gaps", index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
            ws.freeze_panes = "B2"
            for col in ws.columns:
                w = max(len(str(c.value)) if c.value is not None else 0 for c in col[:60])
                ws.column_dimensions[col[0].column_letter].width = min(60, max(10, w + 2))
    print(f"\nWrote {xp} and CSVs in {a.out}/")
    print(summ.to_string(index=False))
    print("\n", PM.T.to_string(header=False))


if __name__ == "__main__":
    main()
