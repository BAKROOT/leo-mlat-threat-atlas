#!/usr/bin/env python3
"""Generate the SATERA LLM-panel prompt pack from the extended catalogue workbook.

Usage (from llm_panel/):  python tools/make_prompts_round1.py <catalogue.xlsx> <outdir>
Writes <outdir>/prompts/ (A, B1-B4, manifest.json, and the small-part variants used in the ChatGPT sensitivity run).

Blind design: threat cards carry only identification, mechanism, vector, asset IDs and
the bound atlas sources (without our evidence level). No STRIDE-tau, goals, tier, class,
evidence level, reach, score or band of the catalogue appears in any prompt.
"""
import json, os, re, sys
import openpyxl

XLSX = sys.argv[1]
OUT = sys.argv[2]
os.makedirs(os.path.join(OUT, "prompts"), exist_ok=True)

wb = openpyxl.load_workbook(XLSX, data_only=True)
# whole numbers are stored as 3 or 3.0 depending on how the workbook was saved; read both as int (in memory only)
for _ws in wb:
    for _row in _ws.iter_rows():
        for _c in _row:
            if isinstance(_c.value, float) and _c.value.is_integer():
                _c.value = int(_c.value)

# ----------------------------------------------------------------- blind scrub
SCRUB = [
    ("A deceptive or deceived aircraft (AC1) broadcasts", "A deceptive or deceived aircraft broadcasts"),
    ("match a false position (AC3).", "match a false position."),
    ("The primary effect is loss of an input; any drift is a secondary tau effect.",
     "The immediate effect is loss of an input; any later clock drift follows from holdover."),
    ("An enabler: it opens the stealth window.",
     "It can widen the margin within which another error goes unflagged."),
    ("The defining tau case: it forges neither a message nor an identity yet violates the integrity of time.",
     "No message is forged and no identity is falsified, yet the time reference ends up wrong."),
    ("Defeated by authenticating the exchange, so Spoofing is primary and tau secondary.",
     "Authenticating the exchange would stop it."),
    ("disclosure is the precondition of injection (T2.09).",
     "what is learned can inform injection (T2.09)."),
    (" An enabler with no consequence of its own.", ""),
    ("Realised by synchronised emitters (AC3).", "Realised by synchronised emitters."),
    (" An enabler of repudiation.", ""),
    ("but corrupts position knowledge, not time, so it stays Tampering.",
     "but corrupts position knowledge, not time."),
    ("A precondition of forged telecommands (T2.11), key abuse (T4.05) and weak-authentication exploitation (T1.12).",
     "It would make forged telecommands (T2.11), key abuse (T4.05) and weak-authentication exploitation (T1.12) easier."),
    ("An enabler: it informs geometry-targeted dropping (T2.04)",
     "It informs geometry-targeted dropping (T2.04)"),
    (" A source in the worked displacement application.", ""),
]
used = {s[0]: 0 for s in SCRUB}

def scrub(text):
    for old, new in SCRUB:
        if old in text:
            text = text.replace(old, new)
            used[old] += 1
    return text

# ----------------------------------------------------------------- data
ws = wb["Catalogue"]
hdr = {ws.cell(4, c).value: c for c in range(1, ws.max_column + 1)}
threats = []
for r in range(5, ws.max_row + 1):
    tid = ws.cell(r, 1).value
    if not tid or not re.match(r"T\d\.\d\d$", str(tid)):
        continue
    threats.append(dict(
        id=tid, layer=ws.cell(r, hdr["Layer"]).value,
        title=ws.cell(r, hdr["Threat"]).value,
        desc=scrub(str(ws.cell(r, hdr["Description (mechanism level)"]).value)),
        vector=str(ws.cell(r, hdr["Attack vector"]).value),
        assets=str(ws.cell(r, hdr["Asset(s)"]).value),
    ))
assert len(threats) == 55, len(threats)
assert all(v > 0 for v in used.values()), [k for k, v in used.items() if v == 0]

wsb = wb["Binding"]
src = {}
for r in range(5, wsb.max_row + 1):
    t = wsb.cell(r, 1).value
    if not t:
        continue
    src.setdefault(t, []).append(dict(
        n=wsb.cell(r, 2).value, rel=wsb.cell(r, 3).value, title=wsb.cell(r, 5).value,
        auth=wsb.cell(r, 6).value, typ=wsb.cell(r, 7).value, peer=wsb.cell(r, 8).value))

wsa = wb["Assets"]
assets = []
for r in range(4, wsa.max_row + 1):
    a = wsa.cell(r, 1).value
    if a and re.match(r"A\d\.\d$", str(a)):
        tb = wsa.cell(r, 7).value
        tb = "none" if tb in (None, "---") else tb
        assets.append((a, wsa.cell(r, 2).value, f"L{int(wsa.cell(r,3).value)}", wsa.cell(r, 4).value, tb))
assert len(assets) == 27

def asset_table():
    return "\n".join(f"{a} | {n} | {l} | {f} | TB: {tb}" for a, n, l, f, tb in assets)

def card(t, vector=True):
    lines = [f"[{t['id']}] Layer {t['layer']} | {t['title']}",
             f"Mechanism: {t['desc']}"]
    if vector:
        lines.append(f"Attack vector: {t['vector']}")
    lines.append(f"Asset(s) named by the authors: {t['assets']}")
    ss = src.get(t["id"], [])
    if ss:
        lines.append("Atlas sources cited by the authors for this threat:")
        for s in ss:
            rel = "path evidence (non-malicious failure)" if s["rel"] == "Path" else "attack evidence"
            lines.append(f"  - #{s['n']} {s['auth']}, \"{s['title']}\" [{s['typ']}; {rel}]")
    else:
        lines.append("Atlas sources cited by the authors for this threat: none cited.")
    return "\n".join(lines)

# ----------------------------------------------------------------- shared text
ROLE = ("You are an independent senior security reviewer with expertise spanning satellite systems, "
        "air-traffic surveillance (ADS-B, multilateration), GNSS and precise timing, and structured threat "
        "modelling (STRIDE, PASTA, DREAD). You are one member of a panel of independent reviewers; you will "
        "not see the others' answers.")

PURPOSE = """PURPOSE AND SCOPE
This is a defensive security review for academic research. The material is a threat model of a civil air-traffic surveillance system, written for peer review in a security journal. The goal is to find errors, gaps and misjudged risks in the threat model so that the system's designers can prioritise defences. You are asked to classify and rate threats at the level of abstraction used in standards and academic papers (STRIDE categories, DREAD-style ratings, public technique names). You are NOT asked for, and must not give, attack procedures, parameters, signal specifications, code, tooling or step-by-step methods; no operational detail is needed for this task. Everything you need is in this prompt."""

RULES_COMMON = """RULES
1. Work closed-book. Do not browse the web, run code, open files or call any tool. Use only this prompt and your own knowledge.
2. Be independent and critical. The authors' text may contain mistakes, duplicates, mis-scoped or invalid threats. Do not assume anything is correct because it is written here. Disagreeing with the wording is expected when you have a reason.
3. Do not invent sources or facts. If you do not know a published source or are unsure of a fact, say so, lower your confidence, and do not pretend to have read it (in rating prompts, list the cited sources you could not assess in "unknown_sources").
4. Stay at the level of a defensive review: classification, rating and reasoning about risk. Do not add attack parameters, procedures, code or tooling.
5. Return exactly ONE JSON object inside ONE code block (```json ... ```). Nothing before it and nothing after it. Valid JSON only: double quotes, no comments, no trailing commas, no line breaks inside strings.
6. Keep every free-text field within its word limit. If you run out of space, shorten the text, never the list of threats."""

SYSTEM_CONTEXT = f"""SYSTEM UNDER REVIEW (LEO-MLAT-B, SATERA project)
A GNSS-independent air-traffic surveillance system. Aircraft 1090 MHz ADS-B (1090ES) emissions are received by a constellation of low-Earth-orbit satellites that measure time of arrival (TOA), frequency of arrival (FOA) and angle of arrival (AOA). Reports are relayed over inter-satellite links and downlinks to a ground Central Processing Station, where a TDOA/FDOA localisation solver and tracker compute each aircraft position, an integrity monitor tests each solution, and an E-MLAT/ADS-B cross-check compares the independent position with the broadcast one and sets an integrity indicator. Results are delivered to air-traffic management (ATM) as ASTERIX CAT020/021 reports. The ADS-B content is an unauthenticated self-report.

LAYERS
L1 signal, payload and timing (aircraft emission, RF front end, payload processor, onboard clocks, GNSS). L2 links and reports (observable reports, inter-satellite links, downlink, backhaul, PTP time transfer, telecommands). L3 central processing (association, ephemeris database, solver and tracker, integrity monitor, cross-check, ASTERIX encoder, hosting platform). L4 operations, governance and ATM consumption (synchronisation policy, orbit determination, key management, configuration and update management, data feed to ATM, the ATM system and controller, personnel and suppliers, operations centre).

ASSET REGISTER (27 assets): ID | asset | layer | function | trust boundary
{asset_table()}

TRUST BOUNDARIES
TB0 system perimeter: air interface (emissions, GNSS), delivery to ATM, organisational interface to personnel and suppliers. TB1 payload and platform (L1) / links (L2). TB2 links (L2) / Central Processing Station (L3). TB3 processing (L3) / governance (L4).

TIMING DEPENDENCY AND EFFECT-TRANSFER CHANNELS
Onboard clocks (A1.5) are disciplined by GNSS (A1.2) and by PTP time-transfer messages (A2.3). A clock offset delta becomes a range bias eps = c*delta on every observable it stamps (channel Phi, L1 to L2, TB1). Bias becomes position error e_h ~ h*eps, with h the HDOP (channel M, L2 to L3, TB2); ephemeris error enters here. Within L3, the monitor and cross-check flag a solution iff eps >= eps_det(h, gamma) (channel V). Error and verdict become a consequence at ATM (channel P, L3 to L4, TB3): displacement happens iff e_h > e_req and no flag is raised. Governance acts back on lower layers by setting a threshold, ephemeris, key or configuration (channel C, L4 to L1-3). Horizontal accuracy requirement e_req = 350 m; c = 0.2998 m/ns; covert clock step delta_cov = 100 ns; so h* = e_req / (c * delta_cov) = 11.67.

ASSUMPTIONS OF THE ANALYSIS
W1 A covert bias step of the size reported in the literature escapes the L1 clock checks. W2 In a STRONG pass (valid solution and HDOP <= h*) the L3 checks detect any single-satellite bias large enough to breach the requirement (no stealth window). W3 In a WEAK pass (valid solution and HDOP > h*) the stealth window is non-empty and no check detects a bias inside it. W4 Observable reports carry no message authentication.

DEFINITIONS
STRIDE-tau categories (choose one PRIMARY, optionally up to two secondary):
S Spoofing: fabricating an emission or observation attributable to a source. T Tampering: biasing, modifying or replay-corrupting an existing observable, data store or configuration. R Repudiation: acting without reliable attribution. I Information disclosure: keys, configuration, precise ephemerides or non-public traffic revealed. D Denial of service: suppressing or degrading the measurements or the geometric observability. E Elevation of privilege: gaining control of a platform, host or command path. tau Temporal integrity: the primary effect is an offset, drift or delay in the time reference of an asset on the timing dependency. Authentication-invariance test: assume every message is authenticated with sound cryptography; if the threat still achieves a time offset, it is tau.
Goals (a threat may serve several): Displacement: a real aircraft is shown at a wrong position with a nominal integrity annotation. Denial: a target, or coverage of a region, is lost or falls below the requirement. Deception: a fabricated target, or a wrong integrity verdict (a missed discrepancy masking a deceptive aircraft, or a false one discrediting a genuine aircraft). Disclosure: keys, configuration, precise ephemerides or non-public traffic revealed; matters as the precondition of a later step.
Adversary classes: AC1 deceptive or deceived aircraft (tier T1); AC2 RF-local emitter (T1); AC3 synchronised multi-emitter (T2); AC4 timing adversary (T2-T3); AC5 link adversary (T2-T3); AC6 ground-segment intruder (T2); AC7 insider or supplier with authorised access (T2).
Tiers: T1p individual with commercial SDRs and public tools; T1c individual with custom software (cost of the order of 10^3 EUR); T2 organised group (several time-synchronised transmitters, published ephemerides, remote IT access); T3 state-level actor (tracks satellites, can place an insider, sustains operations); T3x tier T3 with an insider or two adversary classes.
Evidence levels (for the strongest published support of the threat): E3 observed in operation or documented incident; E2 demonstrated on real hardware; E1 simulated or proof of concept; E0 theoretical or survey; P postulated, no published mechanism (retained, not discarded). Analogue only = the evidence comes from another system (terrestrial multilateration, a ground GNSS receiver, another satellite network) and not from LEO ADS-B multilateration.
Enabler: a threat that has no consequence on a target report by itself and only makes another threat easier or harder to detect."""

RUBRIC = """PA-DREAD RUBRIC (score each dimension as an integer 1-5; you do not compute the total)
D Damage: the consequence class reached at L4. 5 displacement, 4 deception, 3 regional denial or flagged degradation (affected users >= 3), 2 local denial (affected users <= 2), 1 disclosure only. An enabler has no Damage: use null.
R Reproducibility: highest evidence level of the published support. P 1, E0 2, E1 3, E2 4, E3 5; subtract 1 if the evidence is from an analogue system only (minimum 1). Non-malicious failures cited as path evidence do not raise the level.
E Exploitability: tier of the adversary class needed. 5 T1p, 4 T1c, 3 T2, 2 T3, 1 T3x.
A Affected users: extent of the root event. 5 whole constellation, 4 several footprints, 3 one footprint, 2 several aircraft, 1 one aircraft.
Di Discoverability: first layer at which a detection predicate is true. 1 at L1, 2 at L2, 3 at L3, 4 only at L4, 5 no check detects it (5 is the riskier end: hard to discover). Give it twice: Di_strong under the STRONG pass (W2) and Di_weak under the WEAK pass (W3).
Bands of the total S = D+R+E+A+Di (computed by the authors, not by you): Low 5-9, Medium 10-14, High 15-19, Critical 20-25."""

B_SCHEMA = """OUTPUT FORMAT (one JSON object; keys exactly as below. In the template, "a | b | c" means choose exactly one value; text after // is explanation only and must NOT appear in your output)
{
  "meta": {"model_self_report": "<your model name and version as you know it>", "prompt_id": "<PROMPT_ID>", "n_ratings": <integer>},
  "ratings": [
    {
      "id": "<threat id exactly as in the card>",
      "validity": "valid | duplicate | not_a_threat | out_of_scope",
      "related_id": "<id of the duplicate or overlapping threat, or null>",
      "layer": "L1 | L2 | L3 | L4",                       // layer where the threat is realised (entry layer)
      "asset_ids": ["A1.1"],                              // asset IDs from the register that you judge the threat acts on
      "trust_boundary": ["TB0"],                          // boundary or boundaries it sits at or crosses; [] if none
      "stride_primary": "S | T | R | I | D | E | tau",
      "stride_secondary": [],                             // 0-2 values from the same list
      "goals": ["Displacement"],                          // subset of Displacement, Denial, Deception, Disclosure; [] for an enabler
      "is_enabler": false,
      "adversary_class": ["AC2"],                         // 1-2 values from AC1..AC7
      "tier": "T1p | T1c | T2 | T3 | T3x",
      "evidence": "P | E0 | E1 | E2 | E3",
      "analogue_only": false,
      "unknown_sources": [],                              // cited atlas source numbers you could not assess, e.g. [59]
      "D": 4,                                             // 1-5, or null for an enabler
      "R": 3, "E": 4, "A": 2,                             // 1-5
      "Di_strong": 3, "Di_weak": 5,                       // 1-5
      "confidence": 3,                                    // 1 = guess, 5 = certain, for the whole record
      "rationale": "<max 45 words: why these categories and scores>",
      "issue": "<max 30 words: any defect in the threat (wording, scope, duplicate, wrong layer or asset), or empty string>"
    }
  ],
  "batch_comments": {
    "missing_threats": [{"title": "<short>", "layer": "L1", "asset_ids": ["A1.3"], "why": "<max 25 words>"}],   // at most 5; threats you expect for these assets that are not in the batch
    "rubric_ambiguities": ["<max 25 words each>"],        // at most 5; places where the definitions above did not let you decide
    "overall": "<max 60 words: your overall view of this batch>"
  }
}
Format example only (a dummy threat, do not reuse its values): {"id": "T0.00", "validity": "valid", "related_id": null, "layer": "L1", "asset_ids": ["A1.1"], "trust_boundary": ["TB0"], "stride_primary": "S", "stride_secondary": [], "goals": ["Deception"], "is_enabler": false, "adversary_class": ["AC2"], "tier": "T1p", "evidence": "E1", "analogue_only": false, "unknown_sources": [], "D": 4, "R": 3, "E": 5, "A": 2, "Di_strong": 3, "Di_weak": 3, "confidence": 3, "rationale": "...", "issue": ""}"""

A_SCHEMA = """OUTPUT FORMAT (one JSON object; keys exactly as below. In the template, "a | b | c" means choose exactly one value; text after // is explanation only and must NOT appear in your output)
{
  "meta": {"model_self_report": "<your model name and version as you know it>", "prompt_id": "A", "n_threats": <integer>},
  "threats": [
    {
      "title": "<max 10 words>",
      "layer": "L1 | L2 | L3 | L4",
      "asset_ids": ["A1.1"],                              // from the register
      "mechanism": "<max 35 words, high level only>",
      "stride_primary": "S | T | R | I | D | E | tau",
      "goals": ["Displacement"],                          // subset of Displacement, Denial, Deception, Disclosure; [] for an enabler
      "tier": "T1p | T1c | T2 | T3 | T3x",
      "evidence_support": "P | E0 | E1 | E2 | E3",       // strongest published support you know of
      "severity": 3,                                      // your own overall rating, 1 (minor) to 5 (critical)
      "commonly_overlooked": false                        // true if a typical threat model of such a system would miss it
    }
  ],
  "architecture_weaknesses": ["<max 30 words each>"],     // at most 8 weaknesses of the architecture or of its assumptions W1-W4
  "overall": "<max 60 words>"
}"""

# ----------------------------------------------------------------- Prompt A
A = f"""=== PROMPT A: INDEPENDENT THREAT ELICITATION ===
{ROLE}

{PURPOSE}

TASK
Below is the architecture and the vocabulary of a threat model. No threat list is given. From the architecture alone, list the threats you would expect a thorough threat model of this system to contain. Cover all four layers and all 27 assets (each asset should be the target of at least one threat unless you judge it not attackable). Give at most 45 threats, ordered by your own severity, and avoid near-duplicates. Then list up to 8 weaknesses of the architecture or of its stated assumptions.

{RULES_COMMON}

{SYSTEM_CONTEXT}

{A_SCHEMA}

Now produce the JSON object."""

with open(os.path.join(OUT, "prompts", "A_elicitation.txt"), "w", encoding="utf-8") as f:
    f.write(A)

# ----------------------------------------------------------------- Prompts B1-B4
manifest = {"A": {"expected_ids": None, "kind": "elicitation"}}
names = {"L1": ("B1", "B1_L1.txt"), "L2": ("B2", "B2_L2.txt"), "L3": ("B3", "B3_L3.txt"), "L4": ("B4", "B4_L4.txt")}
for layer, (pid, fn) in names.items():
    ts = [t for t in threats if t["layer"] == layer]
    cards = "\n\n".join(card(t) for t in ts)
    ids = [t["id"] for t in ts]
    body = f"""=== PROMPT {pid}: INDEPENDENT BLIND RATING, LAYER {layer} ({len(ts)} threats: {ids[0]} to {ids[-1]}) ===
{ROLE}

{PURPOSE}

TASK
Below are {len(ts)} threat records of a threat model, each with the mechanism, the attack vector, the assets the authors bound it to, and the published sources they cite. The authors' own categories, goals, tier, evidence level and scores are deliberately withheld. For EACH record, give your own independent classification and your own PA-DREAD scores using only the definitions and rubric below. Also judge whether the record is a valid, distinct threat. Return one rating object per record, in the same order, every id exactly once ({len(ts)} objects).

{RULES_COMMON}

{SYSTEM_CONTEXT}

{RUBRIC}

{B_SCHEMA.replace("<PROMPT_ID>", pid)}

THREAT RECORDS ({len(ts)})

{cards}

Now produce the JSON object with {len(ts)} ratings."""
    with open(os.path.join(OUT, "prompts", fn), "w", encoding="utf-8") as f:
        f.write(body)
    manifest[pid] = {"expected_ids": ids, "kind": "rating", "file": fn}


# ----------------------------------------------------------------- small sub-batches (for chat apps whose filter rejects long prompts)
SMALL = 4

def build_b(sid, header, ts, vec):
    ids = [t["id"] for t in ts]
    cards = "\n\n".join(card(t, vec) for t in ts)
    what = "the mechanism, the attack vector, the assets" if vec else "the mechanism, the assets"
    return f"""=== PROMPT {sid}: {header} ===
{ROLE}

{PURPOSE}

TASK
Below are {len(ts)} threat records of a threat model, each with {what} the authors bound it to, and the published sources they cite. The authors' own categories, goals, tier, evidence level and scores are deliberately withheld. For EACH record, give your own independent classification and your own PA-DREAD scores using only the definitions and rubric below. Also judge whether the record is a valid, distinct threat. Return one rating object per record, in the same order, every id exactly once ({len(ts)} objects).

{RULES_COMMON}

{SYSTEM_CONTEXT}

{RUBRIC}

{B_SCHEMA.replace("<PROMPT_ID>", sid)}

THREAT RECORDS ({len(ts)})

{cards}

Now produce the JSON object with {len(ts)} ratings."""

for VDIR, VEC in (("small", True), ("small_nv", False)):
    os.makedirs(os.path.join(OUT, "prompts", VDIR), exist_ok=True)
    for layer, (pid, fn) in names.items():
        ts_all = [t for t in threats if t["layer"] == layer]
        k = -(-len(ts_all) // SMALL)
        size = -(-len(ts_all) // k)
        for j in range(k):
            ts = ts_all[j * size:(j + 1) * size]
            sid = pid + "abcdefgh"[j]
            ids = [t["id"] for t in ts]
            hdr = f"INDEPENDENT BLIND RATING, LAYER {layer}, PART {j+1} OF {k} ({len(ts)} threats: {', '.join(ids)})"
            with open(os.path.join(OUT, "prompts", VDIR, f"{sid}_{layer}.txt"), "w", encoding="utf-8") as f:
                f.write(build_b(sid, hdr, ts, VEC))
            if VEC:
                manifest[sid] = {"expected_ids": ids, "kind": "rating", "file": f"small/{sid}_{layer}.txt", "small": True}

with open(os.path.join(OUT, "prompts", "manifest.json"), "w") as f:
    json.dump(manifest, f, indent=1)

for fn in sorted(os.listdir(os.path.join(OUT, "prompts"))) + ["small/" + x for x in sorted(os.listdir(os.path.join(OUT, "prompts", "small")))][:2]:
    p = os.path.join(OUT, "prompts", fn)
    if os.path.isdir(p):
        continue
    s = open(p, encoding="utf-8").read()
    print(f"{fn:22s} {len(s):7d} chars  ~{len(s)//4:6d} tokens  {len(s.split()):6d} words")
